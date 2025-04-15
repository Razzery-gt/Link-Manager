import webbrowser
import json
import os
import getpass
import bcrypt
import validators
import requests
import csv
import yaml
import xml.etree.ElementTree as ET
from docx import Document
import pandas as pd
import tkinter as tk
from tkinter import filedialog
import pyperclip
from datetime import datetime
from colorama import init, Fore
import re

# Инициализация библиотеки colorama
init(autoreset=True)

# Константы
DOCUMENTS_DIR = os.path.join(os.path.expanduser("~"), "Documents", "LinkManager files")
LINKS_FILENAME = os.path.join(DOCUMENTS_DIR, 'url_links.json')
SETTINGS_FILENAME = os.path.join(DOCUMENTS_DIR, 'settings.json')

# Начальные настройки
default_settings = {
    "password": bcrypt.hashpw("1234".encode(), bcrypt.gensalt()).decode(),
    "password_required": False,
    "show_links": True
}

# Стандартные ссылки
default_links = {
    "открыть_браузер": {
        "url": "https://example.com",
        "date_added": str(datetime.now()),
        "category": "общие",
        "status": "не проверено"
    },
    "гугл": {
        "url": "https://www.google.com",
        "date_added": str(datetime.now()),
        "category": "поисковики",
        "status": "не проверено"
    },
    "яндекс": {
        "url": "https://www.yandex.ru",
        "date_added": str(datetime.now()),
        "category": "поисковики",
        "status": "не проверено"
    }
}

def load_links():
    if os.path.exists(LINKS_FILENAME):
        try:
            with open(LINKS_FILENAME, 'r', encoding='utf-8') as f:
                links = json.load(f)
                for key in links:
                    links[key].setdefault('category', 'без категории')
                    links[key].setdefault('status', 'не проверено')
                return links
        except (json.JSONDecodeError, IOError) as e:
            print(Fore.RED + f"Ошибка загрузки ссылок: {e}.")
    return default_links.copy()

def save_links(links):
    try:
        with open(LINKS_FILENAME, 'w', encoding='utf-8') as f:
            json.dump(links, f, ensure_ascii=False, indent=4)
    except IOError as e:
        print(Fore.RED + f"Ошибка при сохранении ссылок: {e}.")

def load_settings():
    if os.path.exists(SETTINGS_FILENAME):
        try:
            with open(SETTINGS_FILENAME, 'r', encoding='utf-8') as f:
                return json.load(f)
        except (json.JSONDecodeError, IOError) as e:
            print(Fore.RED + f"Ошибка загрузки настроек: {e}.")
    return default_settings.copy()

def save_settings(settings):
    try:
        with open(SETTINGS_FILENAME, 'w', encoding='utf-8') as f:
            json.dump(settings, f, ensure_ascii=False, indent=4)
    except IOError as e:
        print(Fore.RED + f"Ошибка при сохранении настроек: {e}.")

def open_browser(url):
    try:
        webbrowser.open(url)
        print(Fore.GREEN + f"Открываем: {url}")
    except Exception as e:
        print(Fore.RED + f"Ошибка при открытии браузера: {e}")

def is_valid_url(url):
    return validators.url(url)

def check_url_accessibility(url):
    try:
        response = requests.get(url, timeout=5)
        return response.status_code == 200
    except requests.RequestException:
        return False

def show_available_keys(links):
    if links:
        print(Fore.CYAN + "Доступные ключи для открытия браузера:")
        for index, (key, data) in enumerate(links.items(), 1):
            print(f"{index}. {Fore.YELLOW}{key} - {data['url']}")
            print(f"   Категория: {data['category']}")
            print(f"   Статус: {data['status']}\n")
    else:
        print(Fore.RED + "Нет доступных ключей.")

def reset_program():
    if os.path.exists(LINKS_FILENAME):
        os.remove(LINKS_FILENAME)
    if os.path.exists(SETTINGS_FILENAME):
        os.remove(SETTINGS_FILENAME)
    print("Link Manager сброшен к настройкам по умолчанию.")

def hash_password(password):
    return bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()

def verify_password(stored_password, provided_password):
    return bcrypt.checkpw(provided_password.encode(), stored_password.encode())

def count_links(links):
    print(Fore.GREEN + f"Количество сохраненных ссылок: {len(links)}")

def copy_to_clipboard(text):
    try:
        pyperclip.copy(text)
        print(Fore.GREEN + "Ссылка скопирована в буфер обмена.")
    except pyperclip.PyperclipException:
        print(Fore.RED + "Ошибка при копировании в буфер обмена.")

def choose_file(save=False, filetypes=(("JSON files", "*.json"), ("All files", "*.*"))):
    root = tk.Tk()
    root.withdraw()
    if save:
        filepath = filedialog.asksaveasfilename(initialdir=DOCUMENTS_DIR, title="Выберите место для сохранения", filetypes=filetypes)
    else:
        filepath = filedialog.askopenfilename(initialdir=DOCUMENTS_DIR, title="Выберите файл для импорта", filetypes=filetypes)
    return filepath

def export_links(links, filename, format):
    try:
        if format == 'csv':
            with open(filename, mode='w', newline='', encoding='utf-8') as csv_file:
                writer = csv.writer(csv_file)
                writer.writerow(['Key', 'URL', 'Date Added', 'Category', 'Status'])
                for key, data in links.items():
                    writer.writerow([key, data['url'], data['date_added'], data['category'], data['status']])
        elif format == 'json':
            with open(filename, 'w', encoding='utf-8') as f:
                json.dump(links, f, ensure_ascii=False, indent=4)
        elif format == 'yaml':
            with open(filename, 'w', encoding='utf-8') as f:
                yaml.dump(links, f, allow_unicode=True, indent=4)
        elif format == 'xml':
            root = ET.Element("links")
            for key, data in links.items():
                link = ET.SubElement(root, "link")
                ET.SubElement(link, "key").text = key
                ET.SubElement(link, "url").text = data['url']
                ET.SubElement(link, "date_added").text = data['date_added']
                ET.SubElement(link, "category").text = data['category']
                ET.SubElement(link, "status").text = data['status']
            tree = ET.ElementTree(root)
            tree.write(filename, encoding='utf-8', xml_declaration=True)
        elif format == 'docx':
            document = Document()
            document.add_heading('Links', level=1)
            for key, data in links.items():
                document.add_paragraph(f"Key: {key}")
                document.add_paragraph(f"URL: {data['url']}")
                document.add_paragraph(f"Date Added: {data['date_added']}")
                document.add_paragraph(f"Category: {data['category']}")
                document.add_paragraph(f"Status: {data['status']}")
                document.add_paragraph()
            document.save(filename)
        elif format == 'txt':
            with open(filename, 'w', encoding='utf-8') as f:
                for key, data in links.items():
                    f.write(f"Key: {key}\nURL: {data['url']}\nDate Added: {data['date_added']}\nCategory: {data['category']}\nStatus: {data['status']}\n\n")
        elif format == 'xlsx':
            data = []
            for key, link_data in links.items():
                data.append([key, link_data['url'], link_data['date_added'], link_data['category'], link_data['status']])
            df = pd.DataFrame(data, columns=['Key', 'URL', 'Date Added', 'Category', 'Status'])
            df.to_excel(filename, index=False)
        else:
            print(Fore.RED + "Неподдерживаемый формат файла.")
            return
        print(Fore.GREEN + f"Ссылки экспортированы в {filename} в формате {format.upper()}.")
    except Exception as e:
        print(Fore.RED + f"Ошибка при экспорте ссылок: {e}.")

def import_links(filename, format):
    try:
        new_links = {}
        if format == 'csv':
            with open(filename, mode='r', encoding='utf-8') as csv_file:
                reader = csv.DictReader(csv_file)
                for row in reader:
                    key = row['Key']
                    new_links[key] = {
                        'url': row['URL'],
                        'date_added': row['Date Added'],
                        'category': row.get('Category', 'без категории'),
                        'status': row.get('Status', 'не проверено')
                    }
        elif format == 'json':
            with open(filename, 'r', encoding='utf-8') as f:
                new_links = json.load(f)
        elif format == 'yaml':
            with open(filename, 'r', encoding='utf-8') as f:
                new_links = yaml.safe_load(f)
        elif format == 'xml':
            tree = ET.parse(filename)
            root = tree.getroot()
            for link in root.findall('link'):
                key = link.find('key').text
                new_links[key] = {
                    'url': link.find('url').text,
                    'date_added': link.find('date_added').text,
                    'category': link.find('category').text if link.find('category') is not None else 'без категории',
                    'status': link.find('status').text if link.find('status') is not None else 'не проверено'
                }
        elif format == 'docx':
            document = Document(filename)
            key = url = date_added = category = status = None
            for para in document.paragraphs:
                text = para.text.strip()
                if text.startswith("Key: "):
                    key = text[5:]
                elif text.startswith("URL: "):
                    url = text[5:]
                elif text.startswith("Date Added: "):
                    date_added = text[12:]
                elif text.startswith("Category: "):
                    category = text[10:]
                elif text.startswith("Status: "):
                    status = text[8:]
                if key and url and date_added:
                    new_links[key] = {
                        'url': url,
                        'date_added': date_added,
                        'category': category or 'без категории',
                        'status': status or 'не проверено'
                    }
                    key = url = date_added = category = status = None
        elif format == 'txt':
            with open(filename, 'r', encoding='utf-8') as f:
                current_entry = {}
                for line in f:
                    line = line.strip()
                    if line.startswith("Key: "):
                        current_entry['key'] = line[5:]
                    elif line.startswith("URL: "):
                        current_entry['url'] = line[5:]
                    elif line.startswith("Date Added: "):
                        current_entry['date_added'] = line[12:]
                    elif line.startswith("Category: "):
                        current_entry['category'] = line[10:]
                    elif line.startswith("Status: "):
                        current_entry['status'] = line[8:]
                    elif line == "":
                        if 'key' in current_entry:
                            new_links[current_entry['key']] = {
                                'url': current_entry['url'],
                                'date_added': current_entry['date_added'],
                                'category': current_entry.get('category', 'без категории'),
                                'status': current_entry.get('status', 'не проверено')
                            }
                            current_entry = {}
        elif format == 'xlsx':
            df = pd.read_excel(filename)
            for _, row in df.iterrows():
                key = str(row['Key'])
                new_links[key] = {
                    'url': row['URL'],
                    'date_added': str(row['Date Added']),
                    'category': row.get('Category', 'без категории'),
                    'status': row.get('Status', 'не проверено')
                }
        else:
            print(Fore.RED + "Неподдерживаемый формат файла.")
            return

        for key, data in new_links.items():
            if key in url_links:
                print(Fore.YELLOW + f"Ключ '{key}' уже существует. Пропускаем.")
            else:
                url_links[key] = data
                print(Fore.GREEN + f"Импортирована ссылка: {key}")
        save_links(url_links)
        print(Fore.GREEN + f"Ссылки импортированы из {filename}.")

    except Exception as e:
        print(Fore.RED + f"Ошибка при импорте ссылок: {e}.")

def check_all_links():
    print(Fore.CYAN + "\nПроверка доступности всех ссылок...")
    total = len(url_links)
    current = 0
    for key in url_links:
        current += 1
        url = url_links[key]['url']
        print(f"[{current}/{total}] Проверка {key}...")
        try:
            if check_url_accessibility(url):
                url_links[key]['status'] = "🟢 Доступна"
                print(Fore.GREEN + "Успешно")
            else:
                url_links[key]['status'] = "🔴 Недоступна"
                print(Fore.RED + "Недоступна")
        except Exception as e:
            url_links[key]['status'] = "⚠️ Ошибка"
            print(Fore.YELLOW + f"Ошибка: {str(e)}")
    save_links(url_links)
    print(Fore.GREEN + "Проверка завершена!")

def advanced_search():
    print("\nРасширенный поиск:")
    print("1. По ключевому слову")
    print("2. По категории")
    print("3. По статусу")
    print("4. По регулярному выражению")
    print("5. Назад")
    
    choice = menu_option("Выберите тип поиска: ", range(1, 6))
    
    found = {}
    if choice == 1:
        query = input("Введите ключевое слово: ").lower()
        found = {k: v for k, v in url_links.items() if query in k.lower() or query in v['url'].lower()}
    elif choice == 2:
        category = input("Введите категорию: ").lower()
        found = {k: v for k, v in url_links.items() if v.get('category', '').lower() == category}
    elif choice == 3:
        status = input("Введите статус: ").lower()
        found = {k: v for k, v in url_links.items() if v.get('status', '').lower().startswith(status)}
    elif choice == 4:
        try:
            pattern = input("Введите регулярное выражение: ")
            regex = re.compile(pattern, re.IGNORECASE)
            found = {k: v for k, v in url_links.items() if regex.search(k) or regex.search(v['url'])}
        except re.error as e:
            print(Fore.RED + f"Ошибка в регулярном выражении: {e}")
            return
    else:
        return
    
    if found:
        show_available_keys(found)
    else:
        print(Fore.RED + "Совпадений не найдено.")

def menu_option(prompt, options):
    while True:
        try:
            choice = int(input(prompt))
            if choice in options:
                return choice
            else:
                print(Fore.RED + "Неверный ввод. Пожалуйста, попробуйте снова.")
        except ValueError:
            print(Fore.RED + "Неверный ввод. Введите число.")

# Основной цикл
url_links = load_links()
settings = load_settings()

if settings["password_required"]:
    password_attempts = 3
    while password_attempts > 0:
        password_input = getpass.getpass("Введите пароль: ")
        if verify_password(settings["password"], password_input):
            break
        else:
            password_attempts -= 1
            print(Fore.RED + f"Неверно. Осталось попыток: {password_attempts}")
    else:
        print(Fore.RED + "Доступ запрещен.")
        exit()

print(Fore.GREEN + "Добро пожаловать в Link Manager!")
print(Fore.GREEN + "Версия: 1.5")

while True:
    print("\n" + "="*30)
    print("1. Открыть ссылку по ключу")
    print("2. Добавить/Изменить ссылку")
    print("3. Удалить ссылку")
    print("4. Переименовать ссылку")
    print("5. Показать все ссылки")
    print("6. Настройки")
    print("7. Статистика по использованию")
    print("8. Экспорт ссылок")
    print("9. Импорт ссылок")
    print("10. Проверить все ссылки")
    print("11. Поиск ссылок")
    print("12. Выход")

    choice = menu_option("Выберите действие: ", range(1, 13))

    if choice == 1:
        if settings["show_links"]:
            show_available_keys(url_links)
        user_input = input("Введите ключ или номер (или 'copy'): ").strip()
        if user_input.lower() == 'copy':
            key = input("Введите ключ для копирования: ")
            if key in url_links:
                copy_to_clipboard(url_links[key]['url'])
            else:
                print(Fore.RED + "Ключ не найден.")
        elif user_input.isdigit():
            index = int(user_input) - 1
            keys = list(url_links.keys())
            if 0 <= index < len(keys):
                key = keys[index]
                open_browser(url_links[key]['url'])
            else:
                print(Fore.RED + "Неверный номер.")
        elif user_input in url_links:
            open_browser(url_links[user_input]['url'])
        else:
            print(Fore.RED + "Ключ не найден.")

    elif choice == 2:
        key = input("Ключ: ")
        url = input("URL: ")
        if not is_valid_url(url):
            print(Fore.RED + "Некорректный URL!")
            continue
        category = input("Категория (Enter для пропуска): ") or 'без категории'
        url_links[key] = {
            'url': url,
            'date_added': str(datetime.now()),
            'category': category,
            'status': 'не проверено'
        }
        save_links(url_links)
        print(Fore.GREEN + "Ссылка сохранена!")

    elif choice == 3:
        key = input("Введите ключ для удаления: ")
        if key in url_links:
            del url_links[key]
            save_links(url_links)
            print(Fore.GREEN + "Ссылка удалена.")
        else:
            print(Fore.RED + "Ключ не найден.")

    elif choice == 4:
        old_key = input("Текущий ключ: ")
        new_key = input("Новый ключ: ")
        if old_key in url_links:
            url_links[new_key] = url_links.pop(old_key)
            save_links(url_links)
            print(Fore.GREEN + "Ключ изменен.")
        else:
            print(Fore.RED + "Ключ не найден.")

    elif choice == 5:
        show_available_keys(url_links)

    elif choice == 6:
        while True:
            print("\nНастройки:")
            print("1. Изменить пароль")
            print("2. Защита паролем: " + ("ВКЛ" if settings["password_required"] else "ВЫКЛ"))
            print("3. Показ ссылок: " + ("ВКЛ" if settings["show_links"] else "ВЫКЛ"))
            print("4. Сброс программы")
            print("5. Назад")
            sub_choice = menu_option("Выберите: ", range(1, 6))
            
            if sub_choice == 1:
                new_pass = getpass.getpass("Новый пароль: ")
                settings["password"] = hash_password(new_pass)
                save_settings(settings)
                print(Fore.GREEN + "Пароль обновлен!")
            elif sub_choice == 2:
                settings["password_required"] = not settings["password_required"]
                save_settings(settings)
                print(Fore.GREEN + f"Защита паролем {'активна' if settings['password_required'] else 'отключена'}.")
            elif sub_choice == 3:
                settings["show_links"] = not settings["show_links"]
                save_settings(settings)
                print(Fore.GREEN + f"Показ ссылок {'включен' if settings['show_links'] else 'выключен'}.")
            elif sub_choice == 4:
                reset_program()
                url_links = load_links()
                settings = load_settings()
                break
            else:
                break

    elif choice == 7:
        count_links(url_links)
        accessible = sum(1 for v in url_links.values() if v['status'] == "🟢 Доступна")
        print(f"Доступных ссылок: {accessible}")
        print(f"Недоступных: {len(url_links) - accessible}")

    elif choice == 8:
        filetypes = [
            ("CSV", "*.csv"), ("JSON", "*.json"), ("YAML", "*.yaml"),
            ("XML", "*.xml"), ("Word", "*.docx"), ("Text", "*.txt"),
            ("Excel", "*.xlsx"), ("Все файлы", "*.*")
        ]
        filename = choose_file(save=True, filetypes=filetypes)
        if filename:
            format = filename.split('.')[-1].lower()
            export_links(url_links, filename, format)

    elif choice == 9:
        filetypes = [
            ("CSV", "*.csv"), ("JSON", "*.json"), ("YAML", "*.yaml"),
            ("XML", "*.xml"), ("Word", "*.docx"), ("Text", "*.txt"),
            ("Excel", "*.xlsx"), ("Все файлы", "*.*")
        ]
        filename = choose_file(filetypes=filetypes)
        if filename:
            format = filename.split('.')[-1].lower()
            import_links(filename, format)

    elif choice == 10:
        check_all_links()

    elif choice == 11:
        advanced_search()

    elif choice == 12:
        print(Fore.GREEN + "Выходим...")
        break

    else:
        print(Fore.RED + "Неверный ввод. Пожалуйста, попробуйте снова.")
